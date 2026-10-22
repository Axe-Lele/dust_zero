"""엑셀 테이블 <-> 게임 JSON 변환.

규약 (tables/*.xlsx)
  - 시트 이름 = 테이블 이름. 파일 이름은 상관없다(정리용으로 도메인별 파일에 나눠 둠).
  - '#'으로 시작하는 시트는 무시한다(설명, 메모, 검산용).
  - 1행 = 컬럼 키, 2행 = '#'로 시작하는 한국어 설명(무시됨), 3행부터 데이터.
  - A열이 '#'으로 시작하는 행은 무시한다(임시로 끄고 싶은 행에 사용).
  - 값이 [ 또는 { 로 시작하면 JSON으로 읽는다(예: config의 starLifeMin).
  - 'x,y > x,y > ...' 형식의 경로 문자열은 map_entrance.path 에서만 쓴다.
"""
import json, os, glob
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "맑은 고딕"
WORKBOOKS = [
    ("01_Config.xlsx", ["config", "enums", "string_ko"]),
    ("02_Unit.xlsx", ["unit", "skill_param", "unit_boost", "unit_levelup", "unit_star", "unit_profile"]),
    ("03_Enemy.xlsx", ["enemy"]),
    ("04_Stage.xlsx", ["area", "power_skill", "stage", "wave", "map", "map_entrance", "map_tile", "star_reward", "tutorial", "guide"]),
    ("05_Reward.xlsx", ["item", "reward_group", "attendance", "condition", "content_unlock", "daily_pool"]),
    ("06_Gacha.xlsx", ["gacha", "gacha_pool", "gear_gacha"]),
    ("09_Gear.xlsx", ["gear", "gear_surv"]),
    ("07_Survivor.xlsx", ["surv_param", "surv_combo", "surv_set", "surv_wave"]),
    ("08_Shop.xlsx", ["shop_product", "pass_level"]),
    ("10_League.xlsx", ["league_unit", "league_param", "league_enemy", "league_tier", "league_round", "league_syn"]),   # 10.18-c 오토배틀러
    ("11_Tower.xlsx", ["tower_floor", "tower_rule", "tower_bonus"]),   # 10.18-g 먼지 타워
    ("12_Coop.xlsx", ["coop_partner"]),   # 10.18-h 협동 청소(봇 파트너)
]

DESC = {
    "league_unit": {"id": "가전 ID", "hp": "체력(★1·성장 0)", "atk": "공격력", "itv": "공격 간격(초)", "range": "사거리(칸, 근접 1)", "spd": "이동 속도(칸/초)", "mana": "스킬까지 마나", "skillKey": "스킬 이름 문구 키", "descKey": "스킬 설명 문구 키({계수}, {계수_pct})"},
    "league_param": {"unitId": "가전 ID", "key": "스킬 계수 이름(radius 칸, dmg 공격력 배율, stun·freeze·dur 초 등)", "value": "값"},
    "league_enemy": {"id": "적 ID(enemy 와 같은 외형)", "hp": "체력(브론즈·1배)", "atk": "공격력", "itv": "공격 간격", "range": "사거리", "spd": "이동 속도", "armor": "받는 피해 감소(0~1)", "fly": "1 = 공중(근접 가전은 못 때림)", "split": "죽으면 나오는 적", "splitN": "그 수"},
    "league_tier": {"tier": "단계", "nameKey": "이름", "hpMul": "적 체력 배율", "atkMul": "적 공격 배율", "firstCrystal": "첫 완주 크리스탈(한 번)", "winCoin": "승리 1번당 코인(보상 판)", "winPart": "승리 1번당 부품(계열마다)", "finishStone": "완주 시 장비 강화석"},
    "league_round": {"round": "라운드", "nameKey": "이름", "squad": "적ID*마릿수, … (먼지 쪽 위 3줄에 자동 배치)", "rMul": "그 라운드 추가 배율", "boss": "보스 라운드 표시"},
    "tower_floor": {"floor": "층", "nameKey": "층 이름(집)", "rule": "층 규칙 ID(tower_rule, 빈칸 = 규칙 없음)", "squad": "적ID*마릿수, … (리그 적 능력 league_enemy)", "slots": "출전 가전 수",
                    "hpMul": "적 체력 배율", "atkMul": "적 공격 배율", "bonus": "보너스 목표 ID(tower_bonus)", "bonusN": "보너스 목표 수치({n})", "bonusCrystal": "보너스 크리스탈(한 번)",
                    "extra": "guest = 빌려줄 가전 ID / guard = 몇 초마다 오는 추가 무리", "crystal": "첫 클리어 크리스탈", "coin": "첫 클리어 코인", "part": "첫 클리어 부품(계열마다)", "stone": "장비 강화석", "shard": "조각 가전 ID", "shardN": "조각 수"},
    "coop_partner": {"id": "파트너 ID", "unit": "얼굴 가전", "nameKey": "이름", "styleKey": "성향 이름", "descKey": "설명", "style": "행동 규칙(rush 소환 몰빵 / greedy 합치기·강화)",
                     "deck": "파트너가 소환하는 가전(쉼표)", "early": "초반 세기 표시(1~3)", "late": "후반 세기 표시(1~3)"},
    "tower_rule": {"id": "규칙 ID(코드의 규칙 이름)", "kind": "a 판 조건 / b 목표가 바뀜 / c 편성 제한", "nameKey": "이름", "descKey": "설명({p1}, {p1_pct}, {p2}, {p2_pct})", "p1": "값 1", "p2": "값 2"},
    "tower_bonus": {"id": "보너스 목표 ID(코드의 판정 이름)", "nameKey": "문구({n} = 층 표 bonusN)"},
    "league_syn": {"id": "효과 ID", "group": "fam1~4 계열 / P 근접 / H 원거리", "need": "필요한 가전 종류 수", "key": "효과(dr 받는 피해 감소, skillDmg, startMana, chill 적 공속 감소, hp, haste)", "value": "값", "nameKey": "이름", "descKey": "설명({v}, {v_pct})"},
    "config": {"key": "설정 이름", "value": "값", "note": "설명"},
    "enums": {"enum": "enum 이름", "name": "항목", "value": "정수값(다른 테이블에서 쓰는 값)"},
    "string_ko": {"key": "문자열 키", "text": "한국어 문구"},
    "unit": {"id": "가전 ID", "order": "도감 순서", "rarity": "등급 1 R / 2 SR / 3 SSR", "family": "계열 1 생활 / 2 주방 / 3 디지털 / 4 계절",
             "tile": "배치 1 바닥(근접) / 2 가구 위", "cost": "배치 비용(1 = 100W)", "heat": "과열 한도(근접만)", "atk": "공격력",
             "itv": "공격 간격(초)", "range": "사거리(칸)", "block": "저지 수", "sp": "스킬 충전(초)", "dur": "스킬 지속(초)",
             "cool": "초당 냉각(선풍기)", "lineCount": "대사 개수", "atkTarget": "평타 대상 1 지상 / 2 공중 / 3 모두",
             "skillTarget": "스킬 대상 1 지상 / 2 공중 / 3 모두", "nameKey": "이름 문자열 키", "shortKey": "짧은 이름 키",
             "roleKey": "역할 키", "skillKey": "스킬 이름 키", "descKey": "스킬 설명 키"},
    "gear": {"id": "장비 ID(1xxx 공용 / 4xxx 전용)", "grade": "등급 1 일반 / 2 고급 / 3 희귀 / 4 전용", "type": "atk 공격 / heat 내구 / sp 충전 / excl 전용", "unitId": "전용장비의 가전 ID",
             "atk": "공격력·스킬 피해 +비율", "heat": "과열 한도 +비율", "sp": "스킬 충전 시간 −비율", "nameKey": "이름 키", "effKey": "스킬 강화 설명 키",
             "sk1Key": "강화할 스킬 계수(skill_param 키, dur = 지속)", "sk1Op": "add 더하기 / mul 곱하기", "sk1Val": "★0 기준 값",
             "sk2Key": "두 번째 계수", "sk2Op": "", "sk2Val": ""},
    "gear_gacha": {"id": "상품 ID", "nameKey": "이름 키", "costItem": "비용 아이템", "costCount": "비용", "pulls": "뽑는 수", "guaranteeGrade": "이 등급 이상 1개 확정(0 없음)",
                   "w1": "1등급 가중치(만분율)", "w2": "2등급", "w3": "3등급", "w4": "4등급(전용, 9종 균등)"},
    "unit_profile": {"unitId": "가전 ID", "art": "일러스트 경로 접두어(비우면 SD로 표시). _full/_face1~3/_turn/_keys.webp", "keyColor": "정보창 대표 색", "bgColor": "정보창 배경 색", "taglineKey": "한 줄 소개 문자열 키", "profileKey": "프로필 문자열 키", "motifKey": "주요 요소 문자열 키"},
    "unit_star": {"star": "돌파 단계(이 단계로 올릴 때)", "shards": "필요 조각", "coin": "필요 코인", "lvCap": "돌파 후 레벨 상한"},
    "unit_boost": {"unitId": "가전 ID", "tier": "강화 단계(2, 3)", "costMul": "강화 비용 = unit.cost × 이 값", "statMul": "그 단계의 능력 배율(공격·스킬 피해·과열 한도·냉각·발전·버프)",
                   "blockAdd": "저지 수 추가", "rangeAdd": "사거리·범위 추가(칸)", "spMul": "스킬 충전 시간 배율"},
    "shop_product": {"id": "상품 ID", "tab": "탭 1 크리스탈 / 2 일일 / 3 패키지", "order": "진열 순서", "payKind": "결제 1 현금(테스트) / 2 크리스탈 / 3 광고",
                     "price": "가격(현금=원, 크리스탈=개수, 광고=0)", "rewardGroup": "지급 보상 그룹(reward_group)", "limitType": "제한 0 없음 / 1 하루 / 2 주 / 3 계정",
                     "limitCount": "제한 횟수", "firstBonusGroup": "첫 구매 보너스 보상 그룹(0 없음)", "special": "idle:시간 / pass:month / pass:noads",
                     "cooldown": "광고 재시청 대기(초)", "unlockCond": "판매 조건(condition, 0 항상)", "badgeKey": "뱃지 문구 키", "nameKey": "이름 키", "descKey": "설명 키"},
    "pass_level": {"level": "패스 레벨", "xp": "도달 누적 경험치", "freeGroup": "무료 줄 보상 그룹", "paidGroup": "유료(프리미엄) 줄 보상 그룹"},
    "attendance": {"day": "출석부 칸(1~7, 7 다음은 1)", "rewardGroup": "보상 그룹", "big": "1이면 큰 칸으로 강조"},
    "skill_param": {"unitId": "가전 ID", "key": "계수 이름(코드가 읽는 키)", "value": "값", "note": "설명"},
    "unit_levelup": {"id": "구간 ID", "lvMin": "시작 레벨", "lvMax": "끝 레벨", "coinBase": "구간 첫 레벨 코인",
                     "coinStep": "레벨당 코인 증가", "partBase": "구간 첫 레벨 부품", "partStep": "레벨당 부품 증가"},
    "enemy": {"id": "적 ID", "nameKey": "이름 키", "hp": "HP(스테이지 배율 전)", "spd": "이동 속도(칸/초)", "dmg": "과열 피해",
              "itv": "공격 간격(초)", "radius": "크기", "weight": "저지 칸 차지(보스 2)", "fly": "비행 1 / 지상 0",
              "look": "외형", "boss": "보스 1", "ai": "이동 성향 1 우회형 / 2 직진형", "armor": "방어(받는 피해 감소 비율)",
              "split": "쓰러질 때 갈라지는 적 ID", "splitN": "갈라지는 수", "summon": "불러내는 적 ID", "summonEvery": "부르는 간격(초)", "summonN": "한 번에 부르는 수"},
    "area": {"id": "구역 ID(스테이지 탭 위쪽 탭)", "nameKey": "구역 이름 키", "shortKey": "탭에 쓰는 짧은 이름 키", "subKey": "챕터 요약 키", "descKey": "설명 키", "chFrom": "첫 챕터", "chTo": "마지막 챕터",
             "color": "대표 색", "bg": "배경 색"},
    "power_skill": {"id": "전력 스킬 storm 과전류 방출 / boost 전력 증폭", "lv": "단계", "unlockCond": "열리는 조건(condition.id)", "unlockStage": "참고: 이 스테이지를 깨면 열림",
                    "cost": "전력 비용(1 = 100W)", "cd": "다시 쓰기까지(초)", "radius": "범위(칸, storm)", "dmg": "피해 × 스테이지 적 HP 배율(storm)", "stun": "기절(초, 보스 1/3)",
                    "dur": "지속(초, boost)", "atkMul": "공격력 배율", "itvMul": "공격 간격 배율(0.8 = 공속 +25%)", "spMul": "스킬 충전 속도 배율"},
    "gear_surv": {"gearId": "전용장비 ID(gear)", "unitId": "가전 ID", "slot": "leader 리더 무기 / comp 동료", "key": "surv_param 키", "op": "add 더하기 / mul 곱하기", "val": "★0 기준 값(★마다 exclStarBonus만큼 커짐)"},
    "stage": {"id": "스테이지 ID", "order": "진행 순서", "chapter": "챕터", "area": "구역(area.id)", "no": "표기(1-1)", "mapId": "맵 ID", "waveGroup": "wave.groupId",
              "boss": "보스 스테이지 1", "hpMul": "적 HP 배율(= 기준 계정 전투력 × 시뮬레이터 보정)", "dmgMul": "적 과열 피해 배율",
              "refLv": "기준 계정 평균 레벨(권장 레벨로 표시)", "refStar": "기준 계정 돌파", "refGear": "기준 계정 공용장비 능력(+비율)", "unlockCond": "condition.id (0 = 없음)",
              "firstRewardGroup": "첫 클리어 reward_group", "idleCoin": "방치 시급 코인", "idlePart": "방치 시급 부품(계열마다)",
              "idleCrystal": "방치 시급 크리스탈", "startPower": "시작 전력(1 = 100W, 0 = config 값)"},
    "wave": {"groupId": "stage.waveGroup", "seq": "순서", "t": "시작 시각(초)", "enemy": "적 ID", "count": "마리 수",
             "gap": "간격(초)", "entrance": "입구 1~3", "hpMul": "이 웨이브 적 HP 배율(stage.hpMul에 곱함). 후반일수록 높게"},
    "map": {"id": "맵 ID", "floor": "바닥 그림(wood / kitchen / bath / bedroom … 맵 ID와 같음)", "nameKey": "이름 키", "goalKey": "목표물 이름 키",
            "gimKey": "기믹 설명 키(없으면 비움)", "descKey": "설명 키", "noDeployFromEntrance": "입구 앞 배치 금지 칸 수",
            "noDeployBeforeGoal": "목표 앞 배치 금지 칸 수"},
    "map_entrance": {"mapId": "맵 ID", "entrance": "입구 번호", "spawnKey": "입구 이름 키", "path": "경로 x,y > x,y ... (마지막 = 목표물)"},
    "map_tile": {"mapId": "맵 ID", "row": "행 0~9", "tiles": "7칸: X 벽 / H 가구 / P 바닥 / F 가스레인지·뜨거운 배관 / W 물기 / R 러그 / O 기름 웅덩이 / E 목표"},
    "star_reward": {"id": "ID", "chapter": "챕터", "needStars": "필요 별", "rewardGroup": "reward_group"},
    "tutorial": {"id": "스텝 ID", "group": "그룹(빈 cond 그룹은 전투 1판에 1그룹씩 순서대로)", "seq": "그룹 안 순서",
                 "cond": "그룹 조건 비움 / boss / map:맵ID / ps:전력스킬ID / over(처음 과열)", "trigger": "뜨는 조건 1 시작 / 2 앞 스텝 직후 / 3 경과 초 / 4 적 첫 등장 / 5 전력 이상 / 6 과열 비율 / 7 스킬 첫 발동 / 8 보스 등장 / 9 첫 누수 / 10 처음 과열",
                 "param": "조건 값(초, 적 ID, 전력, 비율, 가전 ID)", "until": "닫히는 조건 1 확인 / 2 카드 선택 / 3 배치 / 4 초 뒤 자동(10.16-m부터 안 씀)",
                 "untilParam": "닫힘 값(가전 ID / melee / ranged / 초)", "pause": "1 = 떠 있는 동안 전투 멈춤(10.16-m부터 전부 1)",
                 "focus": "강조 card:가전ID·melee·ranged / hud:life / hud:power / btn:speed / btn:storm·boost / unit:heat / unit:over", "tile": "배치 칸 고정 x,y(비우면 자유)", "tileStage": "칸 고정이 적용되는 스테이지 ID", "afterPrev": "1 = 앞 스텝을 끝낸 뒤에만 조건 확인", "textKey": "안내 문구 키({unit} = 대상 가전 이름)"},
    "guide": {"flow": "안내 묶음 first(새 계정) / levelup(1-1을 처음 깬 뒤, 아직 안 봤으면 1-1을 깰 때)", "seq": "순서",
              "target": "밝힐 화면 요소 home.cta / battle.start / result.lobby / tab.roster / tab.home / roster.unit:가전ID / detail.levelup / detail.close",
              "textKey": "안내 문구 키"},
    "item": {"id": "아이템 ID", "nameKey": "이름 키", "family": "부품 계열(재화는 0)", "unitId": "가전 조각이면 그 가전 ID"},
    "surv_param": {"id": "가전·협동기·적·강화 ID", "slot": "leader 리더 무기 / comp 동료 / combo 협동기 / enemy 적 / passive 공용 강화",
                   "key": "계수 이름(코드가 읽는 키)", "value": "값", "note": "설명"},
    "surv_combo": {"id": "협동기 ID", "leader": "리더 가전", "comp": "동료 가전(최대 레벨일 때 열림)", "nameKey": "이름 키", "descKey": "설명 키"},
    "surv_set": {"family": "계열 1 생활 / 2 주방 / 3 디지털 / 4 계절", "need": "필요 인원(리더 포함)", "key": "효과 키", "value": "값", "descKey": "설명 키"},
    "surv_wave": {"id": "ID", "t0": "시작(초)", "t1": "증가가 끝나는 시각(초). 이후엔 끝값으로 계속 나옴", "enemy": "적 ID", "rate0": "시작 시 초당 마릿수", "rate1": "끝날 때 초당 마릿수",
                  "hp0": "시작 HP 배율(enemy.hp 기준)", "hp1": "끝 HP 배율", "curve": "증가 곡선 지수(1 = 직선, 클수록 후반에 몰림)", "burst": "t0에 한꺼번에 둘러싸며 나오는 수(0 = 없음)", "every": "burst 반복 간격(초, 0 = 한 번)"},
    "daily_pool": {"box": "상자 번호", "pts": "오늘의 청소 일지 점수가 이만큼 되면", "crystal": "받는 크리스탈(하루 한 번)"},
    "reward_group": {"groupId": "그룹 ID", "seq": "순서", "kind": "1 아이템 / 2 가전", "refId": "아이템 ID 또는 가전 ID",
                     "min": "최소 수량", "max": "최대 수량", "prob": "확률(10000 = 100%)"},
    "condition": {"id": "조건 ID", "kind": "1 스테이지 클리어", "p1": "파라미터1(스테이지 ID)", "p2": "파라미터2"},
    "content_unlock": {"id": "메뉴", "cond": "condition.id", "lockKey": "잠김 안내 문구 키"},
    "gacha": {"id": "상품 ID", "poolId": "gacha_pool.poolId", "costItem": "비용 아이템", "costCount": "비용",
              "pulls": "뽑기 횟수", "guaranteeRarity": "마지막 칸 확정 등급(0 = 없음)", "unlockCond": "condition.id", "nameKey": "이름 키"},
    "gacha_pool": {"poolId": "풀 ID", "unitId": "가전 ID", "weight": "가중치(풀 합계 대비 확률)"},
}
CONFIG_NOTE = {
    "tutPullUnit": "소환 안내(1-2 클리어 뒤)의 무료 1회에서 확정으로 나오는 가전",
    "leagueLife": "리그 청결도(0이면 탈락)", "leagueRounds": "리그 라운드 수", "leagueStartGold": "리그 시작 코인", "leagueRoundGold": "라운드마다 받는 코인",
    "leagueWinGold": "이긴 라운드 추가 코인", "leagueInterestPer": "이자: 이만큼 모을 때마다 +1", "leagueInterestMax": "이자 최대", "leagueRerollCost": "상점 새로고침 비용",
    "leagueXpCost": "경험치 구매 비용", "leagueXpBuy": "한 번에 사는 경험치", "leagueXpPerRound": "라운드마다 받는 경험치", "leagueStartLv": "시작 레벨(=배치 수)",
    "leagueLvXp": "[레벨별] 이 레벨이 되려면 필요한 경험치(앞 레벨에서)", "leagueOdds": "[레벨별] 상점 확률 [R, SR, SSR] %", "leagueCost": "구매 비용 [R, SR, SSR]",
    "leagueStarMul": "★1·★2·★3 체력·공격 배율", "leagueGrowthMax": "계정 성장(레벨·돌파) 최대 보너스", "leagueGrowthPow": "성장 보너스가 꽉 차는 전투력 증가량(powOf−1)",
    "leagueExclSkill": "전용장비가 있으면 스킬 효과 +(★마다 exclStarBonus 비율로 커짐)", "leagueManaPerAtk": "공격 1번 마나", "leagueManaPerHit": "맞을 때 마나",
    "towerTime": "먼지 타워 한 층 제한 시간(초, 넘기면 실패)", "towerGuestPow": "손님 가전 성장 = 보유 가전 성장 상위 6대 평균 × 이 값", "towerPurHp": "지키기 층 공기청정기 체력(층 hpMul 곱함)", "towerBossEvery": "타워 보스 패턴 간격(초)",
    "poolSurvPerMin": "청소 일지: 대청소 버틴 1분마다 점수", "poolLeagueWin": "청소 일지: 리그 이긴 라운드마다 점수", "poolLeagueClear": "청소 일지: 리그 8라운드 완주 보너스 점수", "poolCarry": "1 = 못 채운 점수를 다음 날 하루만 어제 몫 상자로", "poolCoopWave": "청소 일지: 협동 청소 막은 웨이브마다 점수", "coopWaveGap": "협동 웨이브 간격(초)", "coopHp0": "협동 적 체력 = 이 값 × 진행 스테이지 기준 계정 힘(refLv·refStar·refGear)", "coopHpG": "협동 웨이브마다 적 체력 배율", "coopBossHp": "협동 보스 체력 추가 배율", "coopLoopMax": "협동: 고리 위 먼지가 이만큼이면 끝(보스 5)", "coopLoopWarn": "협동: 경고 시작", "coopCnt0": "협동 웨이브 기본 마릿수", "coopCntG": "협동 웨이브마다 늘어나는 마릿수", "coopCntMax": "협동 웨이브 최대 마릿수", "coopHoldSpeed": "협동: 근접 가전이 붙잡은 먼지 속도 배율", "coopStartPow": "협동 시작 전력", "coopSum0": "협동 소환 첫 값", "coopSumStep": "협동 소환할 때마다 오르는 값", "coopGenMul": "협동: 발전기 회복·풀가동 충전 배율(× ★)", "coopKillPow": "협동 먼지 처치 전력", "coopBossKillPow": "협동 보스 처치 전력", "coopWavePow0": "협동 웨이브 보너스 기본", "coopWavePowStep": "협동 웨이브 보너스 증가", "coopLineMax": "협동: 고리 칸(근접)에 세울 수 있는 내 가전 수", "coopMoveCd": "협동: 옮긴 가전을 다시 옮기기까지(초)", "coopRerollCost": "협동 바꾸기 값(× ★, 1 = 100W)", "coopStarMax": "협동 최고 ★", "coopStarMul": "협동 ★별 능력 배율", "coopUpCost": "협동 종류 강화 단계별 값", "coopUpStep": "협동 종류 강화 단계마다 공격력", "coopCoinPerWave": "협동 막은 웨이브마다 코인", "coopPassXpPerWave": "협동 웨이브마다 시즌 패스 경험치", "coopPassXpMax": "협동 한 판 시즌 패스 경험치 최대",
    "leagueRoundTime": "라운드 제한 시간(초, 넘기면 패배)", "leagueLoseBase": "질 때 기본 청결도 감소(+남은 적 수, 보스 3)", "leagueBench": "대기석 칸",
    "leagueShopSize": "상점 카드 수", "leagueTrialMin": "상점 가전 종류 수(편성에서, 모자라면 다른 보유 → 체험 가전)", "leaguePoolCopies": "가전마다 한 판에 나오는 장수 [R, SR, SSR](사면 줄고 팔면 돌아옴)", "leagueDailyRewards": "리그 하루 보상 판 수", "passXpLeaguePerWin": "리그 승리 1번당 시즌 패스 경험치",
    "deckMax": "편성 최대 수", "lvMax": "최대 레벨", "lvCapBase": "돌파 전(★0) 레벨 상한", "dupeShards": "중복 소환 시 조각(R/SR/SSR)", "shardRecruit": "조각 영입 필요 수(R/SR/SSR)", "survShardPerMin": "대청소 1분당 리더 조각", "survStonePerMin": "대청소 1분당 장비 강화석", "gearMaxLv": "장비 최대 레벨(1·2·3·전용)", "gearStonePerLv": "강화석 비용 계수", "gearCoinPerLv": "코인 비용 계수", "gearStatPerLv": "레벨당 능력 증가", "gearDismantle": "분해 강화석(등급별)", "gearRefund": "분해 시 투자 강화석 환급 비율", "gearFuseNeed": "합성 필요 수", "exclMaxStar": "전용장비 최대 각성", "exclStarBonus": "각성당 스킬 강화 효과 증가", "exclDupeEssence": "각성 초과 중복 → 정수", "gearEssencePerPull": "장비 소환 1회당 정수", "exclExchange": "전용장비 교환 정수", "survShardMax": "대청소 한 판 최대 조각", "starMax": "최대 돌파", "lvBonus": "레벨당 능력치 증가", "starBonus": "돌파당 능력치 증가", "genLvBonus": "발전기 레벨당 전력 회복·풀가동 충전 증가", "genStarBonus": "발전기 돌파당 전력 회복·풀가동 충전 증가", "pcLvBonus": "컴퓨터 레벨당 공격력 버프 증가", "pcStarBonus": "컴퓨터 돌파당 공격력 버프 증가",
    "idleCapHours": "자동 청소 최대 누적 시간", "maxUnits": "전투 동시 배치 최대 수", "tutGuestMax": "튜토리얼 지원 카드 최대 수(안내에 필요한 가전이 편성에 없으면 그 판에만 붙여 줌)", "startPower": "시작 전력(1 = 100W)",
    "regenPerSec": "초당 전력 회복", "redeployCooldown": "철수 후 재배치 대기(초)", "deployCooldown": "같은 가전 연속 배치 대기(초)", "spawnWarnSec": "적 출현 전 입구 경고 시작(초)", "overheatSec": "과열 정지 시간(초)",
    "lifeMax": "청결도", "startRewardGroup": "시작 지급 reward_group", "attendanceRewardGroup": "출석 reward_group",
    "dupeOverflowRewardGroup": "돌파 최대 이후 중복 전환 reward_group", "repRarityAutoSwap": "더 높은 등급 획득 시 대표 자동 교체(1/0)",
    "starLifeMin": "별 1/2/3개에 필요한 남은 청결도", "queueHeatRatio": "줄 선 적이 가전에 주는 과열 비율",
    "powerMax": "전력 최대치(1 = 100W)", "retreatRefund": "철수 시 환급 비율(배치+강화 비용 기준)", "multitapCost": "멀티탭(배치 칸 +1) 비용 목록, 1 = 100W", "adWatchSec": "보상형 광고 길이(초, 테스트용)", "passSeasonStart": "패스 1시즌 시작일", "passSeasonDays": "패스 시즌 길이(일)", "passXpStageFirst": "패스 경험치: 첫 클리어", "passXpStageRepeat": "패스 경험치: 다시 클리어", "passXpSurvPer30s": "패스 경험치: 대청소 30초마다", "passXpSurvMax": "패스 경험치: 대청소 한 판 최대", "passXpAttend": "패스 경험치: 출석", "passXpAd": "패스 경험치: 광고 1회", "passLevelBuyCrystal": "패스 1레벨 구매 크리스탈", "monthPassDays": "구독 기간(일)", "monthPassDailyGroup": "구독 매일 보상 그룹", "naturalCoolPerSec": "근접 가전 자연 냉각(초당)",
    "fireTileHeatPerSec": "가스레인지 칸 과열(초당)", "survScaleStart": "무한 강화 시작(초)", "survHpGrowthPerMin": "분당 적 HP 증가", "survDmgGrowthPerMin": "분당 적 피해 증가",
    "survSpdGrowthPerMin": "분당 적 이속 증가", "survSpdMax": "적 이속 최대 배율", "survBossFirst": "보스 첫 등장(초)", "survBossEvery": "보스 간격(초)",
    "survBossHp": "보스 HP 배율", "survBossHpStep": "보스 회차당 HP 증가", "survLeaderHp": "대청소 리더 내구도",
    "survMoveSpeed": "대청소 이동 속도(칸/초)", "survPickup": "먼지 조각 흡입 범위(칸)", "survXpBase": "레벨업 경험치 기본", "survXpStep": "레벨당 경험치 증가", "survXpQuad": "레벨 제곱 항(후반 레벨업을 느리게)",
    "survMaxComp": "최대 동료 수", "survCompMaxLv": "동료 최대 레벨", "survWeaponMaxLv": "리더 무기 최대 레벨", "survArena": "방 크기(칸)",
    "survDailyRewards": "하루 보상 판 수", "survPartPer30s": "30초당 부품(계열마다)", "survCoinPerKill": "처치당 코인", "wetTileCoolPerSec": "물기 칸 추가 냉각(초당)", "wetTileEnemySpeedMul": "물기 칸 적 이속 배율",
    "rugTileEnemySpeedMul": "러그 칸 지상 적 이속 배율", "oilTileEnemySpeedMul": "기름 웅덩이 칸 지상 적 이속 배율", "oilTileHeatPerSec": "기름 웅덩이 칸 근접 가전 과열(초당)",
    "armorMax": "적 방어 상한", "stunBossMul": "보스 기절 시간 배율",
}

# ---------- JSON -> 시트 행 ----------
def _rows_for(name, tables):
    t = tables
    if name == "config":
        return ["key", "value", "note"], [{"key": k, "value": v, "note": CONFIG_NOTE.get(k, "")} for k, v in t["config"].items()]
    if name == "enums":
        return ["enum", "name", "value"], [{"enum": e, "name": n, "value": v} for e, m in t["enums"].items() for n, v in m.items()]
    if name == "string_ko":
        return ["key", "text"], [{"key": k, "text": v} for k, v in t["string_ko"].items()]
    if name == "map":
        cols = ["id", "floor", "nameKey", "goalKey", "gimKey", "descKey", "noDeployFromEntrance", "noDeployBeforeGoal"]
        return cols, [{c: m.get(c, "") for c in cols} for m in t["map"]]
    if name == "map_entrance":
        rows = []
        for m in t["map"]:
            for i, (sk, path) in enumerate(zip(m["spawnKeys"], m["paths"])):
                rows.append({"mapId": m["id"], "entrance": i + 1, "spawnKey": sk, "path": " > ".join(f"{x},{y}" for x, y in path)})
        return ["mapId", "entrance", "spawnKey", "path"], rows
    if name == "map_tile":
        return ["mapId", "row", "tiles"], [{"mapId": m["id"], "row": r, "tiles": line} for m in t["map"] for r, line in enumerate(m["tiles"])]
    rows = t[name]
    cols = []
    for r in rows:
        for k in r:
            if k not in cols: cols.append(k)
    return cols, rows

def _cell_value(v):
    if isinstance(v, (list, dict)): return json.dumps(v, ensure_ascii=False)
    return v

def write_workbooks(tables, outdir):
    os.makedirs(outdir, exist_ok=True)
    head_fill = PatternFill("solid", fgColor="1D2226")
    desc_fill = PatternFill("solid", fgColor="E6EAED")
    thin = Side(style="thin", color="C7D0D6")
    for fname, sheets in WORKBOOKS:
        wb = Workbook(); wb.remove(wb.active)
        rd = wb.create_sheet("#설명")
        rd["A1"] = "오퍼레이션: 더스트 제로 테이블"; rd["A1"].font = Font(name=FONT, bold=True, size=14)
        notes = ["시트 이름 = 테이블 이름. '#'으로 시작하는 시트(이 시트 포함)는 변환에서 제외됩니다.",
                 "1행 = 컬럼 키(바꾸지 마세요), 2행 = 설명(무시됨), 3행부터 데이터입니다.",
                 "A열 값 앞에 '#'을 붙이면 그 행은 변환에서 빠집니다. 행을 지우지 않고 잠시 끌 때 쓰세요.",
                 "흰 칸이 입력 칸입니다. 숫자는 숫자로, 문자는 문자로 넣으세요. '1-1' 같은 표기는 날짜로 바뀌지 않게 텍스트 서식이 걸려 있어요.",
                 "수정 후: python tools/build.py  (엑셀 -> data/*.json -> data/tables.js)",
                 "이 파일의 시트: " + ", ".join(sheets)]
        for i, n in enumerate(notes):
            rd.cell(row=3 + i, column=1, value=n).font = Font(name=FONT, size=10)
        rd.column_dimensions["A"].width = 110
        for name in sheets:
            cols, rows = _rows_for(name, tables)
            ws = wb.create_sheet(name)
            desc = DESC.get(name, {})
            for ci, c in enumerate(cols, 1):
                h = ws.cell(row=1, column=ci, value=c)
                h.font = Font(name=FONT, bold=True, color="F3F4F1"); h.fill = head_fill
                d = ws.cell(row=2, column=ci, value=("# " if ci == 1 else "") + desc.get(c, ""))
                d.font = Font(name=FONT, size=9, color="5D6871"); d.fill = desc_fill
                d.alignment = Alignment(wrap_text=True, vertical="top")
            for ri, r in enumerate(rows, 3):
                for ci, c in enumerate(cols, 1):
                    v = _cell_value(r.get(c))
                    cell = ws.cell(row=ri, column=ci, value=v)
                    cell.font = Font(name=FONT, size=10)
                    cell.border = Border(bottom=thin)
                    if isinstance(v, str): cell.number_format = "@"
            for ci, c in enumerate(cols, 1):
                longest = max([len(str(c))] + [len(str(_cell_value(r.get(c)) or "")) for r in rows[:200]])
                ws.column_dimensions[get_column_letter(ci)].width = min(60, max(10, longest * (1.9 if name in ("string_ko",) else 1.2) + 2))
            ws.row_dimensions[2].height = 42
            ws.freeze_panes = "B3"
            ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(2, len(rows) + 2)}"
        wb.save(os.path.join(outdir, fname))

# ---------- 시트 -> JSON ----------
def _norm(v):
    if isinstance(v, str):
        s = v.strip()
        if s[:1] in "[{":
            try: return json.loads(s)
            except Exception: return v
        return v
    if isinstance(v, float) and v.is_integer(): return int(v)
    if hasattr(v, "strftime"):  # 엑셀이 '1-1'을 날짜로 바꿔버린 경우 되돌림
        return f"{v.month}-{v.day}"
    return v

def read_workbooks(indir):
    sheets = {}
    for path in sorted(glob.glob(os.path.join(indir, "*.xlsx"))):
        base = os.path.basename(path)
        if base.startswith("~$") or base.startswith("#"): continue
        wb = load_workbook(path, data_only=True)
        for ws in wb.worksheets:
            if ws.title.startswith("#"): continue
            if ws.title in sheets: raise ValueError(f"시트 '{ws.title}'가 여러 파일에 있습니다: {base}")
            vals = list(ws.iter_rows(values_only=True))
            if not vals: continue
            cols = [c for c in vals[0]]
            rows = []
            for row in vals[1:]:
                if row is None or all(v is None for v in row): continue
                first = row[0]
                if isinstance(first, str) and first.strip().startswith("#"): continue
                rows.append({c: _norm(v) for c, v in zip(cols, row) if c is not None and v is not None})
            sheets[ws.title] = rows
    return assemble(sheets)

def assemble(sh):
    t = {}
    t["config"] = {r["key"]: r.get("value") for r in sh["config"]}
    enums = {}
    for r in sh["enums"]: enums.setdefault(r["enum"], {})[r["name"]] = r["value"]
    t["enums"] = enums
    t["string_ko"] = {r["key"]: r.get("text", "") for r in sh["string_ko"]}
    ents, tiles = {}, {}
    for r in sh["map_entrance"]: ents.setdefault(r["mapId"], []).append(r)
    for r in sh["map_tile"]: tiles.setdefault(r["mapId"], []).append(r)
    maps = []
    for m in sh["map"]:
        m = dict(m); m.setdefault("gimKey", "")
        es = sorted(ents.get(m["id"], []), key=lambda r: r["entrance"])
        m["spawnKeys"] = [e["spawnKey"] for e in es]
        m["paths"] = [[[int(a) for a in p.split(",")] for p in str(e["path"]).split(">")] for e in es]
        m["tiles"] = [str(r["tiles"]) for r in sorted(tiles.get(m["id"], []), key=lambda r: r["row"])]
        maps.append(m)
    t["map"] = maps
    skip = {"config", "enums", "string_ko", "map", "map_entrance", "map_tile"}
    for name, rows in sh.items():
        if name not in skip: t[name] = rows
    return t

"""엑셀 테이블 <-> 게임 JSON 변환.

규약 (tables/*.xlsx)
  - 시트 이름 = 테이블 이름. 파일 이름은 상관없다(정리용으로 도메인별 파일에 나눠 둠).
  - '#'으로 시작하는 시트는 무시한다(설명, 메모, 검산용).
  - 1행 = 컬럼 키, 2행 = '#'로 시작하는 한국어 설명(무시됨), 3행부터 데이터.
  - A열이 '#'으로 시작하는 행은 무시한다(임시로 끄고 싶은 행에 사용).
  - 값이 [ 또는 { 로 시작하면 JSON으로 읽는다(예: config의 starLifeMin).
  - 'x,y > x,y > ...' 형식의 경로 문자열은 map_entrance.path 에서만 쓴다.
"""
import json, os, glob
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "맑은 고딕"
WORKBOOKS = [
    ("01_Config.xlsx", ["config", "enums", "string_ko"]),
    ("02_Unit.xlsx", ["unit", "skill_param", "unit_levelup"]),
    ("03_Enemy.xlsx", ["enemy"]),
    ("04_Stage.xlsx", ["stage", "wave", "map", "map_entrance", "map_tile", "star_reward", "tutorial"]),
    ("05_Reward.xlsx", ["item", "reward_group", "condition", "content_unlock"]),
    ("06_Gacha.xlsx", ["gacha", "gacha_pool"]),
    ("07_Survivor.xlsx", ["surv_param", "surv_combo", "surv_set", "surv_wave"]),
]

DESC = {
    "config": {"key": "설정 이름", "value": "값", "note": "설명"},
    "enums": {"enum": "enum 이름", "name": "항목", "value": "정수값(다른 테이블에서 쓰는 값)"},
    "string_ko": {"key": "문자열 키", "text": "한국어 문구"},
    "unit": {"id": "가전 ID", "order": "도감 순서", "rarity": "등급 1 R / 2 SR / 3 SSR", "family": "계열 1 생활 / 2 주방 / 3 디지털 / 4 계절",
             "tile": "배치 1 바닥(근접) / 2 가구 위", "cost": "배치 비용(1 = 100W)", "heat": "과열 한도(근접만)", "atk": "공격력",
             "itv": "공격 간격(초)", "range": "사거리(칸)", "block": "저지 수", "sp": "스킬 충전(초)", "dur": "스킬 지속(초)",
             "cool": "초당 냉각(선풍기)", "lineCount": "대사 개수", "atkTarget": "평타 대상 1 지상 / 2 공중 / 3 모두",
             "skillTarget": "스킬 대상 1 지상 / 2 공중 / 3 모두", "nameKey": "이름 문자열 키", "shortKey": "짧은 이름 키",
             "roleKey": "역할 키", "skillKey": "스킬 이름 키", "descKey": "스킬 설명 키"},
    "skill_param": {"unitId": "가전 ID", "key": "계수 이름(코드가 읽는 키)", "value": "값", "note": "설명"},
    "unit_levelup": {"id": "구간 ID", "lvMin": "시작 레벨", "lvMax": "끝 레벨", "coinBase": "구간 첫 레벨 코인",
                     "coinStep": "레벨당 코인 증가", "partBase": "구간 첫 레벨 부품", "partStep": "레벨당 부품 증가"},
    "enemy": {"id": "적 ID", "nameKey": "이름 키", "hp": "HP(스테이지 배율 전)", "spd": "이동 속도(칸/초)", "dmg": "과열 피해",
              "itv": "공격 간격(초)", "radius": "크기", "weight": "저지 칸 차지(보스 2)", "fly": "비행 1 / 지상 0",
              "look": "외형", "boss": "보스 1", "ai": "이동 성향 1 우회형 / 2 직진형"},
    "stage": {"id": "스테이지 ID", "order": "진행 순서", "chapter": "챕터", "no": "표기(1-1)", "mapId": "맵 ID", "waveGroup": "wave.groupId",
              "boss": "보스 스테이지 1", "hpMul": "적 HP 배율", "unlockCond": "condition.id (0 = 없음)",
              "firstRewardGroup": "첫 클리어 reward_group", "idleCoin": "방치 시급 코인", "idlePart": "방치 시급 부품(계열마다)",
              "idleCrystal": "방치 시급 크리스탈"},
    "wave": {"groupId": "stage.waveGroup", "seq": "순서", "t": "시작 시각(초)", "enemy": "적 ID", "count": "마리 수",
             "gap": "간격(초)", "entrance": "입구 1~3"},
    "map": {"id": "맵 ID", "floor": "바닥 wood / kitchen / bath", "nameKey": "이름 키", "goalKey": "목표물 이름 키",
            "gimKey": "기믹 설명 키(없으면 비움)", "descKey": "설명 키", "noDeployFromEntrance": "입구 앞 배치 금지 칸 수",
            "noDeployBeforeGoal": "목표 앞 배치 금지 칸 수"},
    "map_entrance": {"mapId": "맵 ID", "entrance": "입구 번호", "spawnKey": "입구 이름 키", "path": "경로 x,y > x,y ... (마지막 = 목표물)"},
    "map_tile": {"mapId": "맵 ID", "row": "행 0~9", "tiles": "7칸: X 벽 / H 가구 / P 바닥 / F 가스레인지 / W 물기 / E 목표"},
    "star_reward": {"id": "ID", "chapter": "챕터", "needStars": "필요 별", "rewardGroup": "reward_group"},
    "tutorial": {"id": "스텝 ID", "group": "그룹(빈 cond 그룹은 전투 1판에 1그룹씩 순서대로)", "seq": "그룹 안 순서",
                 "cond": "그룹 조건 비움 / boss / map:맵ID", "trigger": "뜨는 조건 1 시작 / 2 앞 스텝 직후 / 3 경과 초 / 4 적 첫 등장 / 5 전력 이상 / 6 과열 비율 / 7 스킬 첫 발동 / 8 보스 등장 / 9 첫 누수",
                 "param": "조건 값(초, 적 ID, 전력, 비율, 가전 ID)", "until": "닫히는 조건 1 확인 / 2 카드 선택 / 3 배치 / 4 초 뒤 자동",
                 "untilParam": "닫힘 값(가전 ID / melee / ranged / 초)", "pause": "1 = 떠 있는 동안 전투 멈춤",
                 "focus": "강조 card:가전ID·melee·ranged / hud:life / hud:power / btn:speed / unit:heat", "textKey": "안내 문구 키({unit} = 대상 가전 이름)"},
    "item": {"id": "아이템 ID", "nameKey": "이름 키", "family": "부품 계열(재화는 0)"},
    "surv_param": {"id": "가전·협동기·적·강화 ID", "slot": "leader 리더 무기 / comp 동료 / combo 협동기 / enemy 적 / passive 공용 강화",
                   "key": "계수 이름(코드가 읽는 키)", "value": "값", "note": "설명"},
    "surv_combo": {"id": "협동기 ID", "leader": "리더 가전", "comp": "동료 가전(최대 레벨일 때 열림)", "nameKey": "이름 키", "descKey": "설명 키"},
    "surv_set": {"family": "계열 1 생활 / 2 주방 / 3 디지털 / 4 계절", "need": "필요 인원(리더 포함)", "key": "효과 키", "value": "값", "descKey": "설명 키"},
    "surv_wave": {"id": "ID", "t0": "시작(초)", "t1": "끝(초)", "enemy": "적 ID", "rate0": "시작 시 초당 마릿수", "rate1": "끝날 때 초당 마릿수",
                  "hp0": "시작 HP 배율(enemy.hp 기준)", "hp1": "끝 HP 배율", "curve": "증가 곡선 지수(1 = 직선, 클수록 후반에 몰림)", "burst": "t0에 한꺼번에 둘러싸며 나오는 수(0 = 없음)"},
    "reward_group": {"groupId": "그룹 ID", "seq": "순서", "kind": "1 아이템 / 2 가전", "refId": "아이템 ID 또는 가전 ID",
                     "min": "최소 수량", "max": "최대 수량", "prob": "확률(10000 = 100%)"},
    "condition": {"id": "조건 ID", "kind": "1 스테이지 클리어", "p1": "파라미터1(스테이지 ID)", "p2": "파라미터2"},
    "content_unlock": {"id": "메뉴", "cond": "condition.id", "lockKey": "잠김 안내 문구 키"},
    "gacha": {"id": "상품 ID", "poolId": "gacha_pool.poolId", "costItem": "비용 아이템", "costCount": "비용",
              "pulls": "뽑기 횟수", "guaranteeRarity": "마지막 칸 확정 등급(0 = 없음)", "unlockCond": "condition.id", "nameKey": "이름 키"},
    "gacha_pool": {"poolId": "풀 ID", "unitId": "가전 ID", "weight": "가중치(풀 합계 대비 확률)"},
}
CONFIG_NOTE = {
    "deckMax": "편성 최대 수", "lvMax": "최대 레벨", "starMax": "최대 돌파", "lvBonus": "레벨당 능력치 증가", "starBonus": "돌파당 능력치 증가",
    "idleCapHours": "자동 청소 최대 누적 시간", "maxUnits": "전투 동시 배치 최대 수", "startPower": "시작 전력(1 = 100W)",
    "regenPerSec": "초당 전력 회복", "redeployCooldown": "철수 후 재배치 대기(초)", "overheatSec": "과열 정지 시간(초)",
    "lifeMax": "청결도", "startRewardGroup": "시작 지급 reward_group", "attendanceRewardGroup": "출석 reward_group",
    "dupeOverflowRewardGroup": "돌파 최대 이후 중복 전환 reward_group", "repRarityAutoSwap": "더 높은 등급 획득 시 대표 자동 교체(1/0)",
    "starLifeMin": "별 1/2/3개에 필요한 남은 청결도", "queueHeatRatio": "줄 선 적이 가전에 주는 과열 비율",
    "powerMax": "전력 최대치(1 = 100W)", "retreatRefund": "철수 시 비용 환급 비율", "naturalCoolPerSec": "근접 가전 자연 냉각(초당)",
    "fireTileHeatPerSec": "가스레인지 칸 과열(초당)", "survDuration": "대청소 한 판 길이(초)", "survLeaderHp": "대청소 리더 내구도",
    "survMoveSpeed": "대청소 이동 속도(칸/초)", "survPickup": "먼지 조각 흡입 범위(칸)", "survXpBase": "레벨업 경험치 기본", "survXpStep": "레벨당 경험치 증가",
    "survMaxComp": "최대 동료 수", "survCompMaxLv": "동료 최대 레벨", "survWeaponMaxLv": "리더 무기 최대 레벨", "survArena": "방 크기(칸)",
    "survDailyRewards": "하루 보상 판 수", "survPartPer30s": "30초당 부품(계열마다)", "survCoinPerKill": "처치당 코인", "survClearCrystal": "끝까지 버티면 크리스탈", "wetTileCoolPerSec": "물기 칸 추가 냉각(초당)", "wetTileEnemySpeedMul": "물기 칸 적 이속 배율",
}

# ---------- JSON -> 시트 행 ----------
def _rows_for(name, tables):
    t = tables
    if name == "config":
        return ["key", "value", "note"], [{"key": k, "value": v, "note": CONFIG_NOTE.get(k, "")} for k, v in t["config"].items()]
    if name == "enums":
        return ["enum", "name", "value"], [{"enum": e, "name": n, "value": v} for e, m in t["enums"].items() for n, v in m.items()]
    if name == "string_ko":
        return ["key", "text"], [{"key": k, "text": v} for k, v in t["string_ko"].items()]
    if name == "map":
        cols = ["id", "floor", "nameKey", "goalKey", "gimKey", "descKey", "noDeployFromEntrance", "noDeployBeforeGoal"]
        return cols, [{c: m.get(c, "") for c in cols} for m in t["map"]]
    if name == "map_entrance":
        rows = []
        for m in t["map"]:
            for i, (sk, path) in enumerate(zip(m["spawnKeys"], m["paths"])):
                rows.append({"mapId": m["id"], "entrance": i + 1, "spawnKey": sk, "path": " > ".join(f"{x},{y}" for x, y in path)})
        return ["mapId", "entrance", "spawnKey", "path"], rows
    if name == "map_tile":
        return ["mapId", "row", "tiles"], [{"mapId": m["id"], "row": r, "tiles": line} for m in t["map"] for r, line in enumerate(m["tiles"])]
    rows = t[name]
    cols = []
    for r in rows:
        for k in r:
            if k not in cols: cols.append(k)
    return cols, rows

def _cell_value(v):
    if isinstance(v, (list, dict)): return json.dumps(v, ensure_ascii=False)
    return v

def write_workbooks(tables, outdir):
    os.makedirs(outdir, exist_ok=True)
    head_fill = PatternFill("solid", fgColor="1D2226")
    desc_fill = PatternFill("solid", fgColor="E6EAED")
    thin = Side(style="thin", color="C7D0D6")
    for fname, sheets in WORKBOOKS:
        wb = Workbook(); wb.remove(wb.active)
        rd = wb.create_sheet("#설명")
        rd["A1"] = "오퍼레이션: 더스트 제로 테이블"; rd["A1"].font = Font(name=FONT, bold=True, size=14)
        notes = ["시트 이름 = 테이블 이름. '#'으로 시작하는 시트(이 시트 포함)는 변환에서 제외됩니다.",
                 "1행 = 컬럼 키(바꾸지 마세요), 2행 = 설명(무시됨), 3행부터 데이터입니다.",
                 "A열 값 앞에 '#'을 붙이면 그 행은 변환에서 빠집니다. 행을 지우지 않고 잠시 끌 때 쓰세요.",
                 "흰 칸이 입력 칸입니다. 숫자는 숫자로, 문자는 문자로 넣으세요. '1-1' 같은 표기는 날짜로 바뀌지 않게 텍스트 서식이 걸려 있어요.",
                 "수정 후: python tools/build.py  (엑셀 -> data/*.json -> data/tables.js)",
                 "이 파일의 시트: " + ", ".join(sheets)]
        for i, n in enumerate(notes):
            rd.cell(row=3 + i, column=1, value=n).font = Font(name=FONT, size=10)
        rd.column_dimensions["A"].width = 110
        for name in sheets:
            cols, rows = _rows_for(name, tables)
            ws = wb.create_sheet(name)
            desc = DESC.get(name, {})
            for ci, c in enumerate(cols, 1):
                h = ws.cell(row=1, column=ci, value=c)
                h.font = Font(name=FONT, bold=True, color="F3F4F1"); h.fill = head_fill
                d = ws.cell(row=2, column=ci, value=("# " if ci == 1 else "") + desc.get(c, ""))
                d.font = Font(name=FONT, size=9, color="5D6871"); d.fill = desc_fill
                d.alignment = Alignment(wrap_text=True, vertical="top")
            for ri, r in enumerate(rows, 3):
                for ci, c in enumerate(cols, 1):
                    v = _cell_value(r.get(c))
                    cell = ws.cell(row=ri, column=ci, value=v)
                    cell.font = Font(name=FONT, size=10)
                    cell.border = Border(bottom=thin)
                    if isinstance(v, str): cell.number_format = "@"
            for ci, c in enumerate(cols, 1):
                longest = max([len(str(c))] + [len(str(_cell_value(r.get(c)) or "")) for r in rows[:200]])
                ws.column_dimensions[get_column_letter(ci)].width = min(60, max(10, longest * (1.9 if name in ("string_ko",) else 1.2) + 2))
            ws.row_dimensions[2].height = 42
            ws.freeze_panes = "B3"
            ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(2, len(rows) + 2)}"
        wb.save(os.path.join(outdir, fname))

# ---------- 시트 -> JSON ----------
def _norm(v):
    if isinstance(v, str):
        s = v.strip()
        if s[:1] in "[{":
            try: return json.loads(s)
            except Exception: return v
        return v
    if isinstance(v, float) and v.is_integer(): return int(v)
    if hasattr(v, "strftime"):  # 엑셀이 '1-1'을 날짜로 바꿔버린 경우 되돌림
        return f"{v.month}-{v.day}"
    return v

def read_workbooks(indir):
    sheets = {}
    for path in sorted(glob.glob(os.path.join(indir, "*.xlsx"))):
        base = os.path.basename(path)
        if base.startswith("~$") or base.startswith("#"): continue
        wb = load_workbook(path, data_only=True)
        for ws in wb.worksheets:
            if ws.title.startswith("#"): continue
            if ws.title in sheets: raise ValueError(f"시트 '{ws.title}'가 여러 파일에 있습니다: {base}")
            vals = list(ws.iter_rows(values_only=True))
            if not vals: continue
            cols = [c for c in vals[0]]
            rows = []
            for row in vals[1:]:
                if row is None or all(v is None for v in row): continue
                first = row[0]
                if isinstance(first, str) and first.strip().startswith("#"): continue
                rows.append({c: _norm(v) for c, v in zip(cols, row) if c is not None and v is not None})
            sheets[ws.title] = rows
    return assemble(sheets)

def assemble(sh):
    t = {}
    t["config"] = {r["key"]: r.get("value") for r in sh["config"]}
    enums = {}
    for r in sh["enums"]: enums.setdefault(r["enum"], {})[r["name"]] = r["value"]
    t["enums"] = enums
    t["string_ko"] = {r["key"]: r.get("text", "") for r in sh["string_ko"]}
    ents, tiles = {}, {}
    for r in sh["map_entrance"]: ents.setdefault(r["mapId"], []).append(r)
    for r in sh["map_tile"]: tiles.setdefault(r["mapId"], []).append(r)
    maps = []
    for m in sh["map"]:
        m = dict(m); m.setdefault("gimKey", "")
        es = sorted(ents.get(m["id"], []), key=lambda r: r["entrance"])
        m["spawnKeys"] = [e["spawnKey"] for e in es]
        m["paths"] = [[[int(a) for a in p.split(",")] for p in str(e["path"]).split(">")] for e in es]
        m["tiles"] = [str(r["tiles"]) for r in sorted(tiles.get(m["id"], []), key=lambda r: r["row"])]
        maps.append(m)
    t["map"] = maps
    skip = {"config", "enums", "string_ko", "map", "map_entrance", "map_tile"}
    for name, rows in sh.items():
        if name not in skip: t[name] = rows
    return t
